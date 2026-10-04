import csv
import sys
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
WORKBOOK = ROOT / "corpus_builder/input/work_orders_master.xlsx"
CATALOG = ROOT / "corpus_builder/input/equipment_catalog.csv"
REQUIRED = ["work_order_id", "equipment_id", "manufacturer", "model", "equipment_type",
            "opened_at", "closed_at", "maintenance_type", "reported_symptom", "diagnosis",
            "root_cause", "actions_performed", "technician", "status"]
MAINTENANCE_TYPES = {"corrective", "preventive", "inspection"}
STATUSES = {"closed", "pending"}


def main() -> int:
    if not WORKBOOK.exists():
        print(f"Falta el Excel maestro: {WORKBOOK}")
        return 1
    with CATALOG.open(encoding="utf-8-sig", newline="") as stream:
        catalog = {row["equipment_id"]: row for row in csv.DictReader(stream)}
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    if "work_orders" not in workbook.sheetnames:
        print("Falta la hoja 'work_orders'.")
        return 1
    sheet = workbook["work_orders"]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value).strip() if value is not None else "" for value in next(rows, ())]
    missing = set(REQUIRED) - set(headers)
    if missing:
        print(f"Faltan columnas obligatorias: {', '.join(sorted(missing))}")
        return 1
    indices = {name: headers.index(name) for name in headers}
    errors, seen, count = [], set(), 0
    for row_number, row in enumerate(rows, start=2):
        if not any(value is not None and str(value).strip() for value in row):
            continue
        count += 1
        item = {name: row[index] if index < len(row) else None for name, index in indices.items()}
        for field in REQUIRED:
            if item.get(field) is None or not str(item[field]).strip():
                errors.append(f"Fila {row_number}: falta {field}.")
        order_id = str(item.get("work_order_id", "")).strip()
        if order_id in seen:
            errors.append(f"Fila {row_number}: work_order_id duplicado {order_id}.")
        seen.add(order_id)
        equipment = str(item.get("equipment_id", "")).strip()
        if equipment not in catalog:
            errors.append(f"Fila {row_number}: equipment_id desconocido {equipment}.")
        elif item.get("manufacturer") != catalog[equipment]["manufacturer"] or item.get("model") != catalog[equipment]["model"]:
            errors.append(f"Fila {row_number}: fabricante/modelo no coincide con {equipment}.")
        if str(item.get("maintenance_type", "")).lower() not in MAINTENANCE_TYPES:
            errors.append(f"Fila {row_number}: maintenance_type debe ser {sorted(MAINTENANCE_TYPES)}.")
        if str(item.get("status", "")).lower() not in STATUSES:
            errors.append(f"Fila {row_number}: status debe ser {sorted(STATUSES)}.")
        try:
            opened = date.fromisoformat(str(item["opened_at"])[:10])
            closed_value = item.get("closed_at")
            if closed_value:
                closed = date.fromisoformat(str(closed_value)[:10])
                if closed < opened:
                    errors.append(f"Fila {row_number}: closed_at precede opened_at.")
        except (TypeError, ValueError):
            errors.append(f"Fila {row_number}: fecha inválida, usa YYYY-MM-DD.")
    if count == 0:
        errors.append("El Excel no contiene órdenes.")
    if errors:
        print("Errores de validación:")
        for error in errors:
            print(f"- {error}")
        return 1
    print(f"Validación correcta: {count} órdenes, IDs únicos y equipos/modelos coherentes.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
