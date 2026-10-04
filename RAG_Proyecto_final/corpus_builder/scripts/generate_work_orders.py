import argparse
import csv
from datetime import date, datetime
from html import escape
import sys
from pathlib import Path

from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

ROOT = Path(__file__).resolve().parents[2]
WORKBOOK = ROOT / "corpus_builder/input/work_orders_master.xlsx"
OUTPUT = ROOT / "data/example_corpus/work_orders"


def value(row, indices, name):
    item = row[indices[name]] if indices[name] < len(row) else ""
    if isinstance(item, (date, datetime)):
        return item.strftime("%Y-%m-%d")
    return "" if item is None else str(item)


def main() -> int:
    parser = argparse.ArgumentParser(description="Genera un PDF sintético por fila del Excel maestro.")
    parser.add_argument("--output-dir", type=Path, default=OUTPUT,
                        help="Directorio de salida (por defecto data/example_corpus/work_orders).")
    args = parser.parse_args()
    if not WORKBOOK.exists():
        print(f"No existe {WORKBOOK}. Completa el Excel maestro antes de generar PDFs.")
        return 1
    with CATALOG.open(encoding="utf-8-sig", newline="") as stream:
        catalog = {row["equipment_id"]: row for row in csv.DictReader(stream)}
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    if "work_orders" not in workbook.sheetnames:
        print("Falta la hoja 'work_orders'.")
        return 1
    sheet = workbook["work_orders"]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value).strip() if value is not None else "" for value in next(rows, ())]
    indices = {name: headers.index(name) for name in headers}
    output_dir = args.output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    font_path = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
    bold_font_path = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")
    if font_path.exists() and bold_font_path.exists():
        pdfmetrics.registerFont(TTFont("DejaVu", str(font_path)))
        pdfmetrics.registerFont(TTFont("DejaVu-Bold", str(bold_font_path)))
        pdfmetrics.registerFontFamily("DejaVu", normal="DejaVu", bold="DejaVu-Bold",
                                      italic="DejaVu", boldItalic="DejaVu-Bold")
        styles = getSampleStyleSheet()
        styles["Title"].fontName = "DejaVu-Bold"
        styles["Title"].fontSize = 18
        styles["Title"].leading = 23
        styles["BodyText"].fontName = "DejaVu"
        styles["BodyText"].fontSize = 8.5
        styles["BodyText"].leading = 12
        styles["Italic"].fontName = "DejaVu"
        styles["Italic"].fontSize = 8
    else:
        styles = getSampleStyleSheet()
    generated = 0
    for row in rows:
        if not any(value is not None for value in row):
            continue
        order_id = value(row, indices, "work_order_id")
        equipment_id = value(row, indices, "equipment_id")
        equipment = catalog.get(equipment_id, {})
        title = f"Orden sintética {order_id}"
        story = [Paragraph(title, styles["Title"]),
                 Paragraph("Documento ficticio para demostración académica. No representa una intervención real.", styles["Italic"]),
                 Spacer(1, 0.2 * inch)]
        fields = [
            ("Orden", order_id), ("Equipo", equipment_id),
            ("Tipo de equipo", value(row, indices, "equipment_type")),
            ("Fabricante", value(row, indices, "manufacturer")), ("Modelo", value(row, indices, "model")),
            ("Tipo de mantenimiento", value(row, indices, "maintenance_type")),
            ("Apertura", value(row, indices, "opened_at")), ("Cierre", value(row, indices, "closed_at")),
            ("Estado", value(row, indices, "status")), ("Técnico ficticio", value(row, indices, "technician")),
            ("Síntoma reportado", value(row, indices, "reported_symptom")),
            ("Código de error", value(row, indices, "error_code")),
            ("Diagnóstico", value(row, indices, "diagnosis")), ("Causa raíz", value(row, indices, "root_cause")),
            ("Acciones realizadas", value(row, indices, "actions_performed")),
            ("Refacciones", value(row, indices, "parts_used")),
            ("Tiempo fuera de servicio (h)", value(row, indices, "downtime_hours")),
            ("Recomendaciones", value(row, indices, "recommendations")),
            ("Observaciones", value(row, indices, "notes")),
        ]
        data = [[Paragraph(f"<b>{escape(key)}</b>", styles["BodyText"]), Paragraph(escape(val) if val else "No especificado", styles["BodyText"])]
                for key, val in fields]
        table = Table(data, colWidths=[1.65 * inch, 4.9 * inch], hAlign="LEFT")
        table.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EEF2F6")),
            ("LINEBELOW", (0, 0), (-1, -1), 0.35, colors.HexColor("#CBD5E1")),
            ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7),
            ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))
        story.append(table)
        output = output_dir / f"{order_id}.pdf"
        SimpleDocTemplate(str(output), pagesize=letter, rightMargin=0.65 * inch,
                          leftMargin=0.65 * inch, topMargin=0.65 * inch,
                          bottomMargin=0.65 * inch, title=title).build(story)
        generated += 1
    print(f"Generados {generated} PDFs en {output_dir}. El Excel no se ingirió.")
    return 0


CATALOG = ROOT / "corpus_builder/input/equipment_catalog.csv"

if __name__ == "__main__":
    sys.exit(main())
